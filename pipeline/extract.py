"""Step ① of the Jarvis pipeline: dump a .pptx into raw, loss-free material.

Output (under <case_dir>):
  raw/slides.json   every shape per slide: id, kind, bbox (0-1 of slide), text, table, image ref
  media/            embedded pictures, de-duplicated by content hash
  slides/           full-slide renders (slide-NN.png) via LibreOffice + pdftoppm

Annotation shapes (ovals/rects without text drawn over a picture) are also
projected into that picture's own 0-1 coordinates, so the web app can replay
them as hotspots on the image.

Usage: python pipeline/extract.py cases/case1
"""
import hashlib
import json
import shutil
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SOFFICE_CANDIDATES = [
    "soffice",
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
]
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
MC_NS = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"


def bbox(shape, sw, sh):
    if shape.left is None:
        return None
    return [round(shape.left / sw, 4), round(shape.top / sh, 4),
            round(shape.width / sw, 4), round(shape.height / sh, 4)]


def xml_run_text(r_el):
    """Text of one <a:r>, keeping superscripts/subscripts visible: 10^{-4}, x_{2}."""
    t = r_el.find(f"{A_NS}t")
    text = (t.text or "") if t is not None else ""
    rpr = r_el.find(f"{A_NS}rPr")
    baseline = int(rpr.get("baseline", 0)) if rpr is not None else 0
    if baseline > 0 and text.strip():
        return "^{" + text + "}"
    if baseline < 0 and text.strip():
        return "_{" + text + "}"
    return text


def run_text(r):
    return xml_run_text(r._r)


def xml_cell_text(tc_el):
    """Table cell text with superscripts kept; paragraphs joined by newline."""
    paras = ["".join(xml_run_text(r) for r in p.iterfind(f"{A_NS}r"))
             for p in tc_el.iter(f"{A_NS}p")]
    return "\n".join(p for p in paras if p.strip()).strip()


def is_red(r):
    try:
        rgb = r.font.color.rgb if r.font.color and r.font.color.type is not None else None
    except AttributeError:
        return False
    return rgb is not None and rgb[0] > 0xC0 and rgb[1] < 0x60 and rgb[2] < 0x60


def paragraphs(tf):
    out = []
    for p in tf.paragraphs:
        text = "".join(run_text(r) for r in p.runs).strip()
        if text:
            para = {"text": text, "level": p.level,
                    "bold": any(r.font.bold for r in p.runs)}
            emphasis = [r.text.strip() for r in p.runs if is_red(r) and r.text.strip()]
            if emphasis:
                para["emphasis"] = emphasis  # presenter highlighted these in red
            out.append(para)
    return out


def walk(shapes, sw, sh, media_dir, media_index, group=None):
    for s in shapes:
        rec = {"id": s.shape_id, "name": s.name, "bbox": bbox(s, sw, sh)}
        if group:
            rec["group"] = group
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(s.shapes, sw, sh, media_dir, media_index, group=s.shape_id)
            continue
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
            blob = s.image.blob
            digest = hashlib.sha1(blob).hexdigest()[:12]
            fname = f"{digest}.{s.image.ext}"
            if digest not in media_index:
                (media_dir / fname).write_bytes(blob)
                media_index[digest] = fname
            rec.update(kind="picture", file=fname,
                       px=list(s.image.size),
                       crop=[round(s.crop_left, 4), round(s.crop_top, 4),
                             round(s.crop_right, 4), round(s.crop_bottom, 4)])
        elif getattr(s, "has_table", False) and s.has_table:
            rec.update(kind="table",
                       rows=[[xml_cell_text(c._tc) for c in r.cells] for r in s.table.rows])
        elif getattr(s, "has_chart", False) and s.has_chart:
            ch = s.chart
            rec.update(kind="chart", chart_type=str(ch.chart_type),
                       categories=list(ch.plots[0].categories) if ch.plots else [],
                       series=[{"name": se.name, "values": list(se.values)}
                               for p in ch.plots for se in p.series])
        elif s.shape_type == MSO_SHAPE_TYPE.LINE:
            rec.update(kind="line")
        else:
            paras = paragraphs(s.text_frame) if s.has_text_frame else []
            geom = None
            try:
                geom = str(s.auto_shape_type).split(".")[-1].split(" ")[0]
            except Exception:
                pass
            rec.update(kind="text" if paras else "shape", geom=geom)
            if paras:
                rec["paragraphs"] = paras
        yield rec


def project_annotations(shapes):
    """Map empty shapes (circles/boxes) that sit on top of a picture into picture coords."""
    pics = [s for s in shapes if s["kind"] == "picture" and s["bbox"]]
    for s in shapes:
        if s["kind"] != "shape" or not s["bbox"] or s.get("geom") not in ("OVAL", "RECTANGLE", "ROUNDED_RECTANGLE"):
            continue
        x, y, w, h = s["bbox"]
        cx, cy = x + w / 2, y + h / 2
        for p in pics:
            px, py, pw, ph = p["bbox"]
            if px <= cx <= px + pw and py <= cy <= py + ph and pw * ph > 0.05:
                s["annotates"] = {"picture_id": p["id"],
                                  "rel_bbox": [round((x - px) / pw, 4), round((y - py) / ph, 4),
                                               round(w / pw, 4), round(h / ph, 4)]}
                break


def alternate_content_tables(slide, sw, sh):
    """python-pptx skips shapes wrapped in <mc:AlternateContent>; read their tables directly."""
    out = []
    for frame in slide._element.iterfind(f".//{MC_NS}Choice/{P_NS}graphicFrame"):
        c_nv = frame.find(f"{P_NS}nvGraphicFramePr/{P_NS}cNvPr")
        off = frame.find(f"{P_NS}xfrm/{A_NS}off")
        ext = frame.find(f"{P_NS}xfrm/{A_NS}ext")
        tbl = frame.find(f".//{A_NS}tbl")
        if tbl is None:
            continue
        rows = [[xml_cell_text(tc) for tc in tr.iterfind(f"{A_NS}tc")]
                for tr in tbl.iterfind(f"{A_NS}tr")]
        box = [int(off.get("x")) / sw, int(off.get("y")) / sh, int(ext.get("cx")) / sw, int(ext.get("cy")) / sh]
        out.append({"id": int(c_nv.get("id")), "name": c_nv.get("name"), "kind": "table",
                    "bbox": [round(v, 4) for v in box], "rows": rows, "alternate_content": True})
    return out


def animation_count(slide):
    timing = slide._element.find(f"{P_NS}timing")
    return 0 if timing is None else len(timing.findall(f".//{P_NS}par"))


def render(pptx_path, out_dir):
    soffice = next((c for c in SOFFICE_CANDIDATES if shutil.which(c) or Path(c).exists()), None)
    if not soffice or not shutil.which("pdftoppm"):
        print("skip render: soffice or pdftoppm not found")
        return
    tmp = out_dir / "_render"
    tmp.mkdir(parents=True, exist_ok=True)
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(pptx_path)],
                   check=True, capture_output=True)
    pdf = tmp / (pptx_path.stem + ".pdf")
    for old in out_dir.glob("slide-*.png"):
        old.unlink()
    subprocess.run(["pdftoppm", "-png", "-r", "110", str(pdf), str(out_dir / "slide")], check=True)
    shutil.rmtree(tmp)


def main(case_dir):
    case_dir = Path(case_dir)
    pptx_path = case_dir / "source.pptx"
    prs = Presentation(pptx_path)
    sw, sh = prs.slide_width, prs.slide_height
    media_dir = case_dir / "media"
    media_dir.mkdir(parents=True, exist_ok=True)
    media_index = {}
    slides = []
    for n, slide in enumerate(prs.slides, 1):
        shapes = list(walk(slide.shapes, sw, sh, media_dir, media_index))
        shapes += alternate_content_tables(slide, sw, sh)
        project_annotations(shapes)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        slides.append({"slide": n, "layout": slide.slide_layout.name, "shapes": shapes,
                       "notes": notes, "animation_nodes": animation_count(slide)})
    raw_dir = case_dir / "raw"
    raw_dir.mkdir(exist_ok=True)
    (raw_dir / "slides.json").write_text(
        json.dumps({"source": pptx_path.name, "aspect": round(sw / sh, 4), "slides": slides},
                   ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(slides)} slides, {len(media_index)} unique images")
    render(pptx_path, case_dir / "slides")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "cases/case1")
